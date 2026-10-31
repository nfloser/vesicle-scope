"""Read-only reconciliation of the reviewed DRUM and Dryad mucus snapshots."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
import math
from pathlib import Path

from vesiclescope.validation.dryad_data import read_verified_dryad_v3_source
from vesiclescope.validation import mucus_data


MEASUREMENT_FILE = 'bEV_diffusion_host_response_CVM_1sMSD.csv'
SUMMARY_FILE = 'bEV_diffusion_host_response_CVM_SampleInfo.csv'
SHORT_SPECIES = ('LC', 'LI', 'GV', 'MM')
REPRESENTATION_REL_TOLERANCE = 1e-12


def _csv_rows(raw: bytes, name: str) -> list[list[str]]:
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ValueError(f'{name}: expected UTF-8 CSV') from exc
    try:
        return list(csv.reader(StringIO(text, newline='')))
    except csv.Error as exc:
        raise ValueError(f'{name}: invalid CSV') from exc


def _csv_number(value: str, location: str) -> float:
    try:
        number = float(value)
    except ValueError as exc:
        raise ValueError(f'{location}: expected a numeric measurement') from exc
    if not math.isfinite(number) or number < 0:
        raise ValueError(f'{location}: measurement must be finite and nonnegative')
    return number


def _dryad_measurements(raw: bytes) -> list[list[float]]:
    rows = _csv_rows(raw, MEASUREMENT_FILE)
    width = 80
    expected_samples = [f'Sample {sample}' for sample in range(1, 11) for _ in range(8)]
    expected_species = list(SHORT_SPECIES) * 20
    expected_particles = (['EV'] * 4 + ['WB'] * 4) * 10
    if len(rows) < 4 or rows[0] != expected_samples or rows[1] != expected_species or rows[2] != expected_particles:
        raise ValueError(f'{MEASUREMENT_FILE}: unexpected three-row column header mapping')
    if any(len(row) > width for row in rows[3:]):
        raise ValueError(f'{MEASUREMENT_FILE}: unexpected populated columns')

    columns: list[list[float]] = []
    for column in range(width):
        values = []
        for row_number, row in enumerate(rows[3:], start=4):
            value = row[column] if column < len(row) else ''
            if value == '':
                continue
            values.append(_csv_number(value, f'{MEASUREMENT_FILE} row {row_number}, column {column + 1}'))
        if not values:
            raise ValueError(f'{MEASUREMENT_FILE} column {column + 1}: empty measurement group')
        columns.append(values)
    return columns


def _dryad_summaries(raw: bytes) -> list[list[float]]:
    rows = _csv_rows(raw, SUMMARY_FILE)
    expected = [
        'Sample Number', 'LC bEVs', 'LI bEVs', 'GV bEVs', 'MM bEVs',
        'LC WB', 'LI WB', 'GV WB', 'MM WB', 'pH',
    ]
    if not rows or rows[0] != expected or len(rows) != 11 or any(len(row) != 10 for row in rows):
        raise ValueError(f'{SUMMARY_FILE}: unexpected row or column mapping')
    summaries = []
    for sample, row in enumerate(rows[1:], start=1):
        values = [_csv_number(value, f'{SUMMARY_FILE} sample {sample}') for value in row]
        if values[0] != sample:
            raise ValueError(f'{SUMMARY_FILE}: sample identity does not match row order')
        summaries.append(values[1:9] + [values[0], values[9]])
    return summaries


def _drum_snapshot(path: Path) -> tuple[list[list[float]], list[list[float]], str, int, str]:
    raw = Path(path).read_bytes()
    import hashlib

    digest = hashlib.sha256(raw).hexdigest()
    if digest != mucus_data.REVIEWED_SHA256:
        raise ValueError('workbook SHA-256 does not match the reviewed DRUM original')
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("XLSX reconciliation requires the optional reader: pip install 'vesicle-scope[data]'") from exc

    workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        expected_sheets = ['1s MSD GeoMean'] + [f'Sample {i}' for i in range(1, 11)]
        if workbook.sheetnames != expected_sheets:
            raise ValueError('unexpected sample sheet inventory or order')
        summary_rows = mucus_data._rows(workbook['1s MSD GeoMean'], mucus_data.SUMMARY_HEADERS)
        if len(summary_rows) != 10:
            raise ValueError('summary must contain exactly ten independent samples')
        summaries = []
        columns = []
        for sample, summary in enumerate(summary_rows, start=1):
            parsed_summary = [
                mucus_data._number(value, f'summary row {sample + 1}')
                for value in summary
            ]
            if parsed_summary[8] != sample:
                raise ValueError('summary sample identity does not match the sample sheet')
            summaries.append(parsed_summary)
            sheet = workbook[f'Sample {sample}']
            mucus_data._rows(sheet, mucus_data.PARTICLE_HEADERS)
            rows = list(sheet.iter_rows(min_row=2, max_col=8, values_only=True))
            for column in range(8):
                values = [
                    mucus_data._number(row[column], f'{sheet.title} column {column + 1}')
                    for row in rows if row[column] is not None
                ]
                if not values:
                    raise ValueError(f'{sheet.title} column {column + 1}: empty measurement group')
                columns.append(values)
    finally:
        workbook.close()
    return columns, summaries, digest, len(raw), openpyxl.__version__


def reconcile_mucus_snapshots(workbook_path: Path, dryad_path: Path) -> dict[str, object]:
    """Compare verified snapshots without converting or fitting measurements."""
    raw_files, dryad_container = read_verified_dryad_v3_source(dryad_path)
    dryad_columns = _dryad_measurements(raw_files[MEASUREMENT_FILE])
    dryad_summaries = _dryad_summaries(raw_files[SUMMARY_FILE])
    drum_columns, drum_summaries, drum_digest, drum_size, reader_version = _drum_snapshot(workbook_path)

    group_differences = []
    representation_only = 0
    substantive = 0
    unpaired_drum = 0
    unpaired_dryad = 0
    exact = 0
    for index, (dryad_values, drum_values) in enumerate(zip(dryad_columns, drum_columns)):
        representation = 0
        group_substantive = 0
        group_exact = 0
        for dryad_value, drum_value in zip(dryad_values, drum_values):
            if dryad_value == drum_value:
                group_exact += 1
            elif math.isclose(
                dryad_value,
                drum_value,
                rel_tol=REPRESENTATION_REL_TOLERANCE,
                abs_tol=0.0,
            ):
                representation += 1
            else:
                group_substantive += 1
        group_unpaired_drum = max(0, len(drum_values) - len(dryad_values))
        group_unpaired_dryad = max(0, len(dryad_values) - len(drum_values))
        exact += group_exact
        representation_only += representation
        substantive += group_substantive
        unpaired_drum += group_unpaired_drum
        unpaired_dryad += group_unpaired_dryad
        if representation or group_substantive or group_unpaired_drum or group_unpaired_dryad:
            column = index % 8
            group_differences.append({
                'sample_id': index // 8 + 1,
                'species': mucus_data.SPECIES[column % 4],
                'particle_type': 'bEV' if column < 4 else 'whole_bacterium',
                'dryad_count': len(dryad_values),
                'drum_count': len(drum_values),
                'exact_shared_value_count': group_exact,
                'representation_only_difference_count': representation,
                'substantive_value_mismatch_count': group_substantive,
                'unpaired_drum_value_count': group_unpaired_drum,
                'unpaired_dryad_value_count': group_unpaired_dryad,
            })

    summary_mismatches = sum(
        dryad_value != drum_value
        for dryad_row, drum_row in zip(dryad_summaries, drum_summaries)
        for dryad_value, drum_value in zip(dryad_row, drum_row)
    )
    dryad_count = sum(map(len, dryad_columns))
    drum_count = sum(map(len, drum_columns))
    raw_identical = not (
        representation_only or substantive or unpaired_drum or unpaired_dryad
        or summary_mismatches
    )
    return {
        'schema': 'vesiclescope.mucus-snapshot-reconciliation',
        'version': 1,
        'sources': {
            'dryad': {
                'doi': '10.5061/dryad.cvdncjth8',
                'version_number': 3,
                'license': 'CC0-1.0',
                'container': dryad_container,
            },
            'drum': {
                'doi': '10.13016/kkai-hwng',
                'sha256': drum_digest,
                'size_bytes': drum_size,
                'license': 'CC-BY-NC-ND-3.0-US',
            },
        },
        'reader': {'name': 'openpyxl', 'version': reader_version},
        'comparison_scope': 'ordered one-second MSD particle values and stored sample summaries',
        'group_count': len(drum_columns),
        'dryad_value_count': dryad_count,
        'drum_value_count': drum_count,
        'exact_shared_value_count': exact,
        'representation_relative_tolerance': REPRESENTATION_REL_TOLERANCE,
        'representation_only_difference_count': representation_only,
        'substantive_value_mismatch_count': substantive,
        'unpaired_drum_value_count': unpaired_drum,
        'unpaired_dryad_value_count': unpaired_dryad,
        'summary_cell_mismatch_count': summary_mismatches,
        'raw_snapshots_identical': raw_identical,
        'ready_for_parameter_fitting': False,
        'limitations': [
            'Numeric proximity classifies parser representation only; every such difference is reported.',
            'Unpaired ordered values do not establish why values were excluded or omitted upstream.',
            'Particle values within one sample are not independent biological replicates.',
            'No one-second MSD value is converted into a diffusion coefficient or simulation parameter.',
        ],
        'group_differences': group_differences,
    }
