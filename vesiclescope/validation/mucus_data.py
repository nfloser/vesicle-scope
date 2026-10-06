"""Read-only audit of the exact reviewed DRUM one-second MSD workbook.

This boundary reports measurements, not simulation inputs or fitted parameters.
"""
from __future__ import annotations

import hashlib
from io import BytesIO
import math
from pathlib import Path

REVIEWED_SHA256 = 'df078620f5a9eb8a4005dadf2f179b4fe2196621a7daae0f8119efa0a7c0e957'
SPECIES = ('L. crispatus', 'L. iners', 'G. vaginalis', 'M. mulieris')
SUMMARY_HEADERS = tuple(f'{s} {kind}' for kind in ('bEVs', 'WB') for s in SPECIES) + (
    'Sample Number', 'pH',
)
PARTICLE_HEADERS = tuple(f'{s} {kind}' for kind in ('EV', 'WB') for s in ('LC', 'LI', 'GV', 'MM'))


def _number(value: object, location: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f'{location}: expected a numeric measurement; formulas are not evaluated')
    number = float(value)
    if not math.isfinite(number) or number < 0:
        raise ValueError(f'{location}: measurement must be finite and nonnegative')
    return number


def _header(value: object) -> str:
    return ' '.join(value.split()) if isinstance(value, str) else ''


def _rows(sheet, expected_headers: tuple[str, ...]) -> list[tuple[object, ...]]:
    rows = list(sheet.iter_rows(values_only=True))
    width = len(expected_headers)
    if not rows or tuple(_header(v) for v in rows[0][:width]) != expected_headers:
        raise ValueError(f'{sheet.title}: unexpected column header mapping')
    if any(v is not None for row in rows for v in row[width:]):
        raise ValueError(f'{sheet.title}: unexpected populated columns')
    return [row[:width] for row in rows[1:] if any(v is not None for v in row)]


def audit_mucus_workbook(path: Path) -> dict[str, object]:
    """Verify source identity and summarize all values without removing outliers.

    The original headers omit units. The associated primary Figure 2
    labels MSD in micron squared at one second; no D is inferred.
    """
    raw = Path(path).read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != REVIEWED_SHA256:
        raise ValueError('workbook SHA-256 does not match the reviewed DRUM original')
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("XLSX audit requires the optional reader: pip install 'vesicle-scope[data]'") from exc

    # Parse exactly the hashed bytes, rather than reopening a mutable path.
    workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        expected_sheets = ['1s MSD GeoMean'] + [f'Sample {i}' for i in range(1, 11)]
        if workbook.sheetnames != expected_sheets:
            raise ValueError('unexpected sample sheet inventory or order')
        summaries = _rows(workbook[expected_sheets[0]], SUMMARY_HEADERS)
        if len(summaries) != 10:
            raise ValueError('summary must contain exactly ten independent samples')
        groups = []
        for sample, summary in enumerate(summaries, start=1):
            if _number(summary[8], f'summary I{sample + 1}') != sample:
                raise ValueError('summary sample identity does not match the sample sheet')
            ph = _number(summary[9], f'summary J{sample + 1}')
            if not 0 < ph <= 14:
                raise ValueError('sample pH must be in (0, 14]')
            sheet = workbook[f'Sample {sample}']
            # Retain blank-row positions for precise original cell attribution.
            _rows(sheet, PARTICLE_HEADERS)
            rows = list(sheet.iter_rows(min_row=2, max_col=8, values_only=True))
            for column in range(8):
                letter = chr(ord('A') + column)
                values = [
                    _number(row[column], f'{sheet.title} {letter}{row_index}')
                    for row_index, row in enumerate(rows, start=2)
                    if row[column] is not None
                ]
                if not values:
                    raise ValueError(f'{sheet.title} {letter}: empty measurement group')
                zero_count = sum(v == 0 for v in values)
                # Nonnegative-product convention: a zero produces GM = 0;
                # no replacement, offset or positive-only filtering occurs.
                mean = 0.0 if zero_count else math.exp(math.fsum(math.log(v) for v in values) / len(values))
                reported = _number(summary[column], f'summary {letter}{sample + 1}')
                groups.append({
                    'sample_id': sample,
                    'pH': ph,
                    'pH_group': 'low' if ph <= 4.2 else 'high',
                    'species': SPECIES[column % 4],
                    'particle_type': 'bEV' if column < 4 else 'whole_bacterium',
                    'source_sheet': sheet.title,
                    'source_column': letter,
                    'particle_count': len(values),
                    'zero_count': zero_count,
                    'reported_geometric_mean': reported,
                    'recomputed_geometric_mean': mean,
                    'difference_recomputed_minus_reported': mean - reported,
                })
    finally:
        workbook.close()

    return {
        'schema': 'vesiclescope.measured-mucus-audit',
        'version': 1,
        'source': {
            'doi': '10.13016/kkai-hwng',
            'publication_doi': '10.1038/s41522-025-00866-9',
            'filename': '2025_10_01 Data Figs2.xlsx',
            'sha256': digest,
            'size_bytes': len(raw),
            'license': 'CC-BY-NC-ND-3.0-US',
            'license_url': 'https://creativecommons.org/licenses/by-nc-nd/3.0/us/',
            'attribution': 'Darby Steinman, University of Maryland DRUM, 2025',
        },
        'reader': {'name': 'openpyxl', 'version': openpyxl.__version__},
        'evidence': 'measured_in_related_context',
        'context': 'bacterial extracellular vesicles and whole bacteria in human cervicovaginal mucus',
        'quantity': 'mean_squared_displacement',
        'lag_time_s': 1.0,
        'msd_unit': 'micron^2',
        'unit_status': 'mapped from associated primary Figure 2; not declared in workbook headers',
        'unit_source': 'https://www.nature.com/articles/s41522-025-00866-9/figures/2',
        'unit_transformation': 'none; no scaling or conversion applied',
        'independent_sample_count': 10,
        'zero_policy': 'retain; geometric mean is zero when any observation is zero',
        'exclusion_policy': 'no values removed; upstream exclusions not reconstructed',
        'ready_for_parameter_fitting': False,
        'interpretation': 'local noncommercial measurement audit; not tumor validation, not diffusion estimation',
        'limitations': [
            'Particle values within one sample are not independent biological replicates.',
            'Reported versus recomputed differences do not establish an error without exclusion and rounding provenance.',
            'Single-lag MSD does not establish normal diffusion or reconstruct trajectories.',
            'Do not redistribute modified source data under the original ND license.',
        ],
        'groups': groups,
    }
