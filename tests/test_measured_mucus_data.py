"""Offline acceptance tests use generated synthetic workbooks, never raw data."""
import contextlib
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from vesiclescope.cli import main
from vesiclescope.validation import mucus_data


@unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'optional data reader missing')
class MeasuredMucusDataTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'synthetic.xlsx'
        self.book = Workbook()
        summary = self.book.active
        summary.title = '1s MSD GeoMean'
        summary.append(list(mucus_data.SUMMARY_HEADERS))
        for sample in range(1, 11):
            summary.append([4.0] * 8 + [sample, 4.2 if sample <= 5 else 4.3])
            sheet = self.book.create_sheet(f'Sample {sample}')
            sheet.append(list(mucus_data.PARTICLE_HEADERS))
            sheet.append([1.0] * 8)
            sheet.append([16.0] * 8)
        self.book.save(self.path)

    def audit(self):
        self.book.save(self.path)
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        with patch.object(mucus_data, 'REVIEWED_SHA256', digest):
            return mucus_data.audit_mucus_workbook(self.path)

    def test_exact_source_digest_is_required(self):
        with self.assertRaisesRegex(ValueError, 'SHA-256'):
            mucus_data.audit_mucus_workbook(self.path)

    def test_sample_identity_particle_types_and_geometric_mean(self):
        before = self.path.read_bytes()
        report = self.audit()
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(len(report['groups']), 80)
        self.assertEqual(report['independent_sample_count'], 10)
        self.assertEqual(report['msd_unit'], 'micron^2')
        self.assertFalse(report['ready_for_parameter_fitting'])
        group = report['groups'][0]
        self.assertEqual(group['particle_type'], 'bEV')
        self.assertAlmostEqual(group['recomputed_geometric_mean'], 4.0)
        self.assertEqual(group['particle_count'], 2)
        self.assertEqual(group['pH_group'], 'low')
        self.assertEqual(report['groups'][4]['particle_type'], 'whole_bacterium')
        self.assertEqual(report['groups'][40]['pH_group'], 'high')

    def test_zero_is_counted_not_silently_removed(self):
        self.book['Sample 1']['A2'] = 0
        group = self.audit()['groups'][0]
        self.assertEqual(group['zero_count'], 1)
        self.assertEqual(group['recomputed_geometric_mean'], 0)
        self.assertEqual(group['particle_count'], 2)

    def test_invalid_measurements_and_formula_rejected(self):
        for value in (-1, 'not a measurement', '=1+1', True):
            with self.subTest(value=value):
                self.book['Sample 1']['A2'] = value
                with self.assertRaisesRegex(ValueError, 'Sample 1.*A2'):
                    self.audit()

    def test_empty_group_and_changed_headers_rejected(self):
        self.book['Sample 1']['A2'] = None
        self.book['Sample 1']['A3'] = None
        with self.assertRaisesRegex(ValueError, 'empty'):
            self.audit()
        self.book['Sample 1']['A2'] = 1
        self.book['Sample 1']['A1'] = 'wrong species'
        with self.assertRaisesRegex(ValueError, 'header'):
            self.audit()

    def test_sample_numbers_must_not_be_reordered_or_duplicated(self):
        self.book['1s MSD GeoMean']['I2'] = 2
        with self.assertRaisesRegex(ValueError, 'sample identity'):
            self.audit()

    def test_nonfinite_values_rejected_at_numeric_boundary(self):
        for value in (float('inf'), float('-inf'), float('nan')):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, 'finite'):
                    mucus_data._number(value, 'synthetic measurement')

    def test_unexpected_sheets_columns_and_invalid_ph_rejected(self):
        self.book.create_sheet('unexpected')
        with self.assertRaisesRegex(ValueError, 'inventory'):
            self.audit()
        del self.book['unexpected']
        self.book['Sample 1']['I2'] = 7
        with self.assertRaisesRegex(ValueError, 'populated columns'):
            self.audit()
        self.book['Sample 1']['I2'] = None
        self.book['1s MSD GeoMean']['J2'] = 15
        with self.assertRaisesRegex(ValueError, 'pH'):
            self.audit()

    def test_cli_is_deterministic_and_handles_errors(self):
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        outputs = []
        with patch.object(mucus_data, 'REVIEWED_SHA256', digest):
            for _ in range(2):
                output = io.StringIO()
                with contextlib.redirect_stdout(output):
                    self.assertEqual(main(['data', 'audit-mucus', str(self.path)]), 0)
                outputs.append(output.getvalue())
        self.assertEqual(outputs[0], outputs[1])
        self.assertEqual(json.loads(outputs[0])['schema'], 'vesiclescope.measured-mucus-audit')
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main(['data', 'audit-mucus', str(self.path)]), 2)


if __name__ == '__main__':
    unittest.main()
